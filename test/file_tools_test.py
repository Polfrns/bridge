import importlib.util
import io
import os
from pathlib import Path
import re
import shutil
import stat
import sys
import tarfile
import tempfile
import unittest
import zipfile

spec = importlib.util.spec_from_file_location('file_tools', Path(__file__).parents[1] / 'scripts/file-tools.py')
tool = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = tool
spec.loader.exec_module(tool)


class FileToolsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_text_and_code_extensions_are_read_without_execution(self):
        for extension in tool.TEXT_EXTENSIONS - {'.eml'}:
            with self.subTest(extension=extension):
                path = self.root / ('script' + extension)
                path.write_text("import os; os.remove('DO_NOT_EXECUTE')", encoding='utf-8')
                result = tool.prepare_file(path, self.root)
                self.assertIn('DO_NOT_EXECUTE', result.text)
        path = self.root / 'utf16.txt'
        path.write_text('Texto español', encoding='utf-16')
        self.assertEqual(tool.prepare_file(path, self.root).text, 'Texto español')

    def test_office_previews_include_headers_notes_and_formulas(self):
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from pptx import Presentation
        doc = Document()
        doc.add_paragraph('Body original')
        doc.sections[0].header.paragraphs[0].text = 'Header original'
        path = self.root / 'sample.docx'
        doc.save(path)
        prepared = tool.prepare_file(path, self.root)
        self.assertIn('Header original', prepared.text)
        if shutil.which('soffice'):
            self.assertTrue(prepared.images)
        tool.validate_file(path, self.root)
        doc.add_paragraph('Second change')
        doc.save(path)
        self.assertIn('Body original', tool.prepare_file(path, self.root).text)
        self.assertIn('Second change', tool.prepare_file(path, self.root).text)
        book = Workbook()
        book.active['A1'] = 'Revenue'
        book.active['B1'] = '=SUM(1,2)'
        path = self.root / 'sample.xlsx'
        book.save(path)
        self.assertIn('=SUM(1,2)', tool.prepare_file(path, self.root).text)
        tool.validate_file(path, self.root)
        self.assertEqual(load_workbook(path).active['B1'].value, '=SUM(1,2)')
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text = 'Slide original'
        slide.notes_slide.notes_text_frame.text = 'Speaker note'
        path = self.root / 'sample.pptx'
        ppt.save(path)
        self.assertIn('Speaker note', tool.prepare_file(path, self.root).text)
        tool.validate_file(path, self.root)

    def test_images_normalize_and_original_bytes_are_preserved(self):
        from PIL import Image
        from pillow_heif import register_heif_opener
        register_heif_opener()
        for extension in ['png', 'jpg', 'gif', 'webp', 'bmp', 'tif', 'tiff', 'avif', 'heic', 'heif']:
            with self.subTest(extension=extension):
                path = self.root / ('image.' + extension)
                Image.new('RGB', (20, 16), 'red').save(path)
                before = path.read_bytes()
                result = tool.prepare_file(path, self.root)
                with Image.open(result.images[0]) as preview:
                    self.assertEqual(preview.format, 'PNG')
                self.assertEqual(path.read_bytes(), before)
                tool.validate_file(path, self.root)
        svg = self.root / 'image.svg'
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><rect width="20" height="20" fill="red"/></svg>')
        self.assertTrue(Path(tool.prepare_file(svg, self.root).images[0]).is_file())
        svg.write_text('<svg><image href="file:///etc/passwd"/></svg>')
        with self.assertRaises(tool.FileToolError):
            tool.prepare_file(svg, self.root)

    def test_pdf_text_and_scanned_page_preview(self):
        import pymupdf
        path = self.root / 'sample.pdf'
        doc = pymupdf.open()
        doc.new_page().insert_text((30, 60), 'PDF original')
        doc.new_page()
        doc.save(path)
        doc.close()
        result = tool.prepare_file(path, self.root)
        self.assertIn('PDF original', result.text)
        self.assertEqual(len(result.images), 2)
        tool.validate_file(path, self.root)
        if shutil.which('tesseract'):
            from PIL import Image, ImageDraw, ImageFont
            image = Image.new('RGB', (900, 140), 'white')
            font = ImageFont.truetype('DejaVuSans.ttf', 42)
            ImageDraw.Draw(image).text((20, 30), 'SCANNED ORIGINAL', font=font, fill='black')
            stream = io.BytesIO(); image.save(stream, 'PNG')
            scan = pymupdf.open(); page = scan.new_page(width=900, height=140)
            page.insert_image(page.rect, stream=stream.getvalue())
            scan_path = self.root / 'scan.pdf'; scan.save(scan_path); scan.close()
            self.assertIn('SCANNED ORIGINAL', tool.prepare_file(scan_path, self.root).text)
        path.write_bytes(b'%PDF-1.7 fake')
        with self.assertRaises(Exception):
            tool.validate_file(path, self.root)

    def test_zip_tar_and_7z_are_extracted_with_folders(self):
        import py7zr
        path = self.root / 'project.py'
        path.write_text('print("original")')
        for extension in ['zip', 'tar', 'tar.gz', '7z']:
            archive = self.root / ('project.' + extension)
            if extension == 'zip':
                with zipfile.ZipFile(archive, 'w') as z:
                    z.write(path, 'src/project.py')
            elif extension == '7z':
                with py7zr.SevenZipFile(archive, 'w') as z:
                    z.write(path, 'src/project.py')
            else:
                with tarfile.open(archive, 'w:gz' if extension.endswith('gz') else 'w') as z:
                    z.add(path, 'src/project.py')
            result = tool.prepare_file(archive, self.root)
            self.assertIn('original', result.text)
            self.assertEqual((Path(result.archive_root) / 'src/project.py').read_bytes(), path.read_bytes())
            tool.validate_file(archive, self.root)
        import gzip, bz2, lzma
        for extension, compressor in [('gz', gzip), ('bz2', bz2), ('xz', lzma)]:
            archive = self.root / ('project.py.' + extension)
            with compressor.open(archive, 'wb') as output:
                output.write(path.read_bytes())
            result = tool.prepare_file(archive, self.root)
            self.assertEqual((Path(result.archive_root) / 'project.py').read_bytes(), path.read_bytes())
            tool.validate_file(archive, self.root)

    def test_archive_paths_links_duplicates_and_bombs_are_rejected(self):
        for index, name in enumerate(['../outside', '/absolute', 'C:/outside', 'a\\..\\outside', 'nul.txt', 'a. /x']):
            archive = self.root / f'unsafe-{index}.zip'
            with zipfile.ZipFile(archive, 'w') as z:
                z.writestr(name, 'bad')
            with self.assertRaises(tool.FileToolError):
                tool.prepare_file(archive, self.root)
        archive = self.root / 'link.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            entry = zipfile.ZipInfo('link')
            entry.external_attr = (stat.S_IFLNK | 0o777) << 16
            z.writestr(entry, '/etc/passwd')
        with self.assertRaises(tool.FileToolError):
            tool.prepare_file(archive, self.root)
        archive = self.root / 'duplicate.zip'
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('A.py', 'a')
            z.writestr('a.py', 'b')
        with self.assertRaises(tool.FileToolError):
            tool.prepare_file(archive, self.root)
        with tarfile.open(self.root / 'link.tar', 'w') as z:
            entry = tarfile.TarInfo('link')
            entry.type = tarfile.SYMTYPE
            entry.linkname = '../outside'
            z.addfile(entry)
        with self.assertRaises(tool.FileToolError):
            tool.prepare_file(self.root / 'link.tar', self.root)
        archive = self.root / 'bomb.zip'
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('huge.txt', b'0' * (tool.MAX_FILE_BYTES + 1))
        with self.assertRaises(tool.FileToolError):
            tool.prepare_file(archive, self.root)
        self.assertFalse((self.root.parent / 'outside').exists())

    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice required')
    def test_legacy_and_opendocument_round_trip(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph('Legacy text preserved')
        source = self.root / 'sample.docx'
        doc.save(source)
        for extension in ['doc', 'rtf', 'odt']:
            target = self.root / ('result.' + extension)
            tool.convert_office(source, target, self.root)
            self.assertIn('Legacy text preserved', tool.prepare_file(target, self.root).text)
            tool.validate_file(target, self.root)
        from openpyxl import Workbook
        book = Workbook()
        book.active['A1'] = 'Legacy cell'
        source = self.root / 'sample.xlsx'
        book.save(source)
        for extension in ['xls', 'ods']:
            target = self.root / ('result.' + extension)
            tool.convert_office(source, target, self.root)
            self.assertIn('Legacy cell', tool.prepare_file(target, self.root).text)
            tool.validate_file(target, self.root)
        from pptx import Presentation
        ppt = Presentation()
        ppt.slides.add_slide(ppt.slide_layouts[1]).shapes.title.text = 'Legacy slide'
        source = self.root / 'sample.pptx'
        ppt.save(source)
        for extension in ['ppt', 'odp']:
            target = self.root / ('result.' + extension)
            tool.convert_office(source, target, self.root)
            self.assertIn('Legacy slide', tool.prepare_file(target, self.root).text)
            tool.validate_file(target, self.root)

    def test_structured_text_deliverables_are_parsed_without_execution(self):
        valid = {'ok.json': '{"a": 1}', 'ok.jsonl': '{"a": 1}\n\n{"b": 2}\n', 'ok.csv': 'a,"b, c"\n1,2\n',
                 'ok.tsv': 'a\tb\n', 'ok.py': "import os\nos.remove('DO_NOT_EXECUTE')\n", 'ok.toml': 'a = 1\n',
                 'free.md': '# any text', 'free.html': '<p>unclosed'}
        for name, text in valid.items():
            with self.subTest(name=name):
                (self.root / name).write_text(text, encoding='utf-8')
                tool.validate_file(self.root / name, self.root)
        (self.root / 'legacy.csv').write_bytes('Año;Señal\n'.encode('cp1252'))
        tool.validate_file(self.root / 'legacy.csv', self.root)
        invalid = {'bad.json': '{"a":', 'bad.jsonl': '{"a": 1}\nnot json\n', 'bad.csv': 'a,"unterminated\n',
                   'bad.py': 'def broken(:\n', 'bad.toml': 'a = = 1\n'}
        for name, text in invalid.items():
            with self.subTest(name=name), self.assertRaises((tool.FileToolError, ValueError)):
                (self.root / name).write_text(text, encoding='utf-8')
                tool.validate_file(self.root / name, self.root)
        (self.root / 'bad-utf8.py').write_bytes(b'print("\xff")\n')
        with self.assertRaises(tool.FileToolError):
            tool.validate_file(self.root / 'bad-utf8.py', self.root)
        self.assertFalse((self.root / 'DO_NOT_EXECUTE').exists())

    @unittest.skipUnless(shutil.which('soffice'), 'LibreOffice required')
    def test_recalc_refreshes_stale_values_and_reports_formula_errors(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet['A1'] = 21
        sheet['B1'] = '=A1*2'
        sheet['C1'] = '=1/0'
        source = self.root / 'formulas.xlsx'
        book.save(source)
        # Plant a stale cached result: only a real recalculation replaces it.
        stale = self.root / 'stale.xlsx'
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(stale, 'w', zipfile.ZIP_DEFLATED) as copy:
            for item in original.infolist():
                data = original.read(item)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    data = re.sub(rb'(<f>A1\*2</f>)(<v\s*/>|<v></v>)?', rb'\1<v>999</v>', data)
                    self.assertIn(b'<v>999</v>', data)
                copy.writestr(item, data)
        result = tool.recalculate_workbook(stale, self.root / 'work' / 'checked.xlsx', self.root)
        self.assertEqual(result['error_count'], 1)
        self.assertEqual(result['errors'][0]['cell'], 'C1')
        self.assertEqual(result['errors'][0]['value'], '#DIV/0!')
        from openpyxl import load_workbook
        checked = load_workbook(result['path'], data_only=True)
        self.assertEqual(checked.active['B1'].value, 42)
        with self.assertRaises(tool.FileToolError):
            tool.recalculate_workbook(stale, self.root / 'work' / 'checked.ods', self.root)

    @unittest.skipUnless(shutil.which('ffmpeg') and shutil.which('ffprobe'), 'FFmpeg required')
    def test_media_metadata_and_video_frames(self):
        for extension in ['wav', 'mp3', 'ogg', 'flac', 'm4a', 'aac', 'opus']:
            target = self.root / ('audio.' + extension)
            tool.run_tool(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=1', str(target)])
            self.assertIn('audio', tool.prepare_file(target, self.root).text)
            tool.validate_file(target, self.root)
        for extension in ['mp4', 'mov', 'mkv', 'webm', 'avi']:
            target = self.root / ('video.' + extension)
            tool.run_tool(['ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'color=red:s=32x32:d=1', str(target)])
            result = tool.prepare_file(target, self.root)
            self.assertGreater(len(result.images), 0)
            tool.validate_file(target, self.root)

    def test_reader_and_converter_cannot_access_outside_workspace(self):
        with self.assertRaises(tool.FileToolError):
            tool.workspace_file(self.root, '../outside')
        with self.assertRaises(tool.FileToolError):
            tool.xml_root(b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')


if __name__ == '__main__':
    unittest.main()
